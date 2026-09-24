from __future__ import annotations

import csv
import io
import json
import math
import re
from threading import Lock
from time import monotonic
from contextlib import asynccontextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile, status
from fastapi.encoders import jsonable_encoder
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response, StreamingResponse
from sqlalchemy import func, or_, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from zipfile import BadZipFile, ZipFile

from .config import settings
from .db import (
    AuthSession,
    Company,
    CooperationRecord,
    ImportBatch,
    PerformanceChange,
    PayoutRecord,
    ProjectType,
    ProjectTypeRule,
    SystemMeta,
    User,
    get_db,
    utc_now,
)
from .domain import (
    calculate_current_status,
    calculate_due_amount,
    calculate_paid_total,
    is_supported_rule_combination,
    status_label,
)
from .schemas import (
    ApiMessage,
    CompanyCreate,
    CompanyUpdate,
    DashboardSummary,
    ImportCommitRequest,
    ImportPreview,
    ImportRequest,
    ImportRow,
    LoginRequest,
    MilestoneRequest,
    MilestoneUpdateRequest,
    OverrideRequest,
    PasswordResetRequest,
    PayoutCreate,
    PayoutUpdate,
    PayoutVoidRequest,
    ProjectTypeCreate,
    ProjectTypeUpdate,
    RecordCreate,
    RecordUpdate,
    RuleCreate,
    UserCreate,
    UserPublic,
    UserUpdate,
    public_user,
)
from .security import (
    create_session,
    delete_session,
    get_current_user,
    hash_password,
    require_admin,
    token_hash,
    verify_password,
)
from .time_utils import parse_date_value, today_beijing


MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
LOGIN_WINDOW_SECONDS = 15 * 60
LOGIN_MAX_ATTEMPTS = 5
_login_attempts: Dict[Tuple[str, str], List[float]] = {}
_login_attempts_lock = Lock()


def yuan_to_cents(value: Optional[float]) -> Optional[int]:
    if value is None:
        return None
    return int((Decimal(str(value)) * 100).quantize(Decimal("1")))


def cents_to_yuan(value: Optional[int]) -> Optional[float]:
    if value is None:
        return None
    return float(Decimal(value) / Decimal(100))


def percent_to_bps(value: Optional[float]) -> Optional[int]:
    if value is None:
        return None
    return int((Decimal(str(value)) * 100).quantize(Decimal("1")))


def bps_to_percent(value: Optional[int]) -> Optional[float]:
    if value is None:
        return None
    return float(Decimal(value) / Decimal(100))


def normalize_name(value: str) -> str:
    return re.sub(r"\s+", "", value or "").strip().casefold()


def ensure_system_meta(db: Session) -> None:
    row = db.get(SystemMeta, "data_revision")
    if not row:
        db.add(SystemMeta(key="data_revision", value="0"))
        db.commit()


def bump_revision(db: Session) -> None:
    row = db.get(SystemMeta, "data_revision")
    if not row:
        row = SystemMeta(key="data_revision", value="1")
        db.add(row)
    else:
        row.value = str(int(row.value) + 1)


def bootstrap_admin(db: Session) -> None:
    if db.scalar(select(func.count(User.id))) or not settings.initial_admin_password:
        return
    username = settings.initial_admin_username or "admin"
    if db.scalar(select(User).where(User.username == username)):
        return
    db.add(
        User(
            username=username,
            display_name=settings.initial_admin_name,
            password_hash=hash_password(settings.initial_admin_password),
            role="admin",
            status="active",
        )
    )
    bump_revision(db)
    db.commit()


@asynccontextmanager
async def lifespan(_: FastAPI):
    db = next(get_db())
    try:
        ensure_system_meta(db)
        bootstrap_admin(db)
    finally:
        db.close()
    yield


app = FastAPI(title="个人项目绩效收益记账系统 API", version="0.1.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def require_owned_company(db: Session, company_id: str, user: User) -> Company:
    company = db.get(Company, company_id)
    if not company or company.owner_id != user.id:
        raise HTTPException(status_code=404, detail="公司不存在")
    return company


def require_record(db: Session, record_id: str, user: User) -> CooperationRecord:
    record = db.get(CooperationRecord, record_id)
    if not record or record.owner_id != user.id:
        raise HTTPException(status_code=404, detail="合作记录不存在")
    return record


def get_payouts(db: Session, record_id: str) -> List[PayoutRecord]:
    return list(db.scalars(select(PayoutRecord).where(PayoutRecord.cooperation_record_id == record_id).order_by(PayoutRecord.received_date, PayoutRecord.created_at)))


def rule_for_date(db: Session, project_type_id: str, record_date: date) -> Optional[ProjectTypeRule]:
    return db.scalar(
        select(ProjectTypeRule)
        .where(
            ProjectTypeRule.project_type_id == project_type_id,
            ProjectTypeRule.is_active.is_(True),
            ProjectTypeRule.effective_from <= record_date,
            or_(ProjectTypeRule.effective_to.is_(None), ProjectTypeRule.effective_to >= record_date),
        )
        .order_by(ProjectTypeRule.version_no.desc())
    )


def type_name(db: Session, project_type_id: str) -> str:
    project_type = db.get(ProjectType, project_type_id)
    return project_type.name if project_type else "未知类型"


def company_name(db: Session, company_id: str) -> str:
    company = db.get(Company, company_id)
    return company.name if company else "未知公司"


def record_payload(db: Session, record: CooperationRecord) -> Dict[str, Any]:
    payouts = get_payouts(db, record.id)
    paid_total = calculate_paid_total(payouts)
    due = calculate_due_amount(record)
    current_status = calculate_current_status(record, payouts)
    return {
        "id": record.id,
        "owner_id": record.owner_id,
        "company_id": record.company_id,
        "company_name": company_name(db, record.company_id),
        "project_type_id": record.project_type_id,
        "project_type_name": type_name(db, record.project_type_id),
        "rule_id": record.rule_id,
        "record_date": record.record_date,
        "work_status": record.work_status,
        "completion_date": record.completion_date,
        "publicity_status": record.publicity_status,
        "publicity_date": record.publicity_date,
        "advance_received_date": record.advance_received_date,
        "full_received_date": record.full_received_date,
        "participation_mode": record.participation_mode,
        "my_ratio_percent": bps_to_percent(record.my_ratio_bps),
        "standard_performance_yuan": cents_to_yuan(
            record.override_standard_cents if record.override_standard_cents is not None else record.standard_performance_cents
        ),
        "override_standard_yuan": cents_to_yuan(record.override_standard_cents),
        "manual_due_amount_yuan": cents_to_yuan(record.manual_due_amount_cents),
        "my_due_amount_yuan": cents_to_yuan(due),
        "paid_total_yuan": cents_to_yuan(paid_total),
        "outstanding_yuan": cents_to_yuan(max(due - paid_total, 0)) if due is not None else None,
        "current_status": current_status,
        "current_status_label": status_label(current_status),
        "snapshot_publicity_required": record.snapshot_publicity_required,
        "snapshot_collection_stage": record.snapshot_collection_stage,
        "snapshot_payout_pattern": record.snapshot_payout_pattern,
        "note": record.note,
        "override_reason": record.override_reason,
        "archived_at": record.archived_at,
        "payouts": [payout_payload(item) for item in payouts],
        "created_at": record.created_at,
        "updated_at": record.updated_at,
    }


def payout_payload(payout: PayoutRecord) -> Dict[str, Any]:
    return {
        "id": payout.id,
        "cooperation_record_id": payout.cooperation_record_id,
        "batch_type": payout.batch_type,
        "amount_yuan": cents_to_yuan(payout.amount_cents),
        "received_date": payout.received_date,
        "note": payout.note,
        "is_void": payout.is_void,
    }


def change_payload(field_name: str, old_value: Any, new_value: Any, reason: str, record: CooperationRecord) -> PerformanceChange:
    return PerformanceChange(
        owner_id=record.owner_id,
        cooperation_record_id=record.id,
        field_name=field_name,
        old_value=json.dumps(old_value, ensure_ascii=False, default=str),
        new_value=json.dumps(new_value, ensure_ascii=False, default=str),
        reason=reason,
    )


def login_rate_key(request: Request, username: str) -> Tuple[str, str]:
    client_ip = request.client.host if request.client else "unknown"
    return client_ip, username.strip().casefold()


def check_login_rate_limit(key: Tuple[str, str]) -> None:
    now = monotonic()
    with _login_attempts_lock:
        for stale_key, attempts in list(_login_attempts.items()):
            recent = [attempt for attempt in attempts if now - attempt < LOGIN_WINDOW_SECONDS]
            if recent:
                _login_attempts[stale_key] = recent
            else:
                _login_attempts.pop(stale_key, None)
        attempts = _login_attempts.get(key, [])
        if len(attempts) >= LOGIN_MAX_ATTEMPTS:
            retry_after = max(1, int(LOGIN_WINDOW_SECONDS - (now - attempts[0])) + 1)
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail="登录尝试过多，请稍后再试",
                headers={"Retry-After": str(retry_after)},
            )


def record_login_failure(key: Tuple[str, str]) -> None:
    now = monotonic()
    with _login_attempts_lock:
        attempts = [attempt for attempt in _login_attempts.get(key, []) if now - attempt < LOGIN_WINDOW_SECONDS]
        if len(attempts) >= LOGIN_MAX_ATTEMPTS:
            retry_after = max(1, int(LOGIN_WINDOW_SECONDS - (now - attempts[0])) + 1)
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail="登录尝试过多，请稍后再试",
                headers={"Retry-After": str(retry_after)},
            )
        attempts.append(now)
        _login_attempts[key] = attempts
        if len(_login_attempts) > 5000:
            oldest_key = min(_login_attempts, key=lambda candidate: _login_attempts[candidate][-1])
            _login_attempts.pop(oldest_key, None)


def clear_login_failures(key: Tuple[str, str]) -> None:
    with _login_attempts_lock:
        _login_attempts.pop(key, None)


@app.get("/api/v1/health")
def health() -> Dict[str, str]:
    return {"status": "ok"}


@app.post("/api/v1/auth/login", response_model=UserPublic)
def login(payload: LoginRequest, request: Request, response: Response, db: Session = Depends(get_db)) -> UserPublic:
    rate_key = login_rate_key(request, payload.username)
    check_login_rate_limit(rate_key)
    user = db.scalar(select(User).where(User.username == payload.username.strip()))
    if not user or user.status != "active" or not verify_password(payload.password, user.password_hash):
        record_login_failure(rate_key)
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="用户名或密码错误")
    clear_login_failures(rate_key)
    raw_token = create_session(db, user)
    user.last_login_at = utc_now()
    db.commit()
    response.set_cookie(
        key=settings.cookie_name,
        value=raw_token,
        max_age=settings.session_days * 24 * 60 * 60,
        httponly=True,
        secure=settings.cookie_secure,
        samesite="lax",
        path="/",
    )
    return public_user(user)


@app.post("/api/v1/auth/logout", response_model=ApiMessage)
def logout(request: Request, response: Response, db: Session = Depends(get_db)) -> ApiMessage:
    delete_session(db, request.cookies.get(settings.cookie_name))
    db.commit()
    response.delete_cookie(settings.cookie_name, path="/")
    return ApiMessage(message="已退出登录")


@app.get("/api/v1/auth/me", response_model=UserPublic)
def me(user: User = Depends(get_current_user)) -> UserPublic:
    return public_user(user)


@app.get("/api/v1/admin/users", response_model=List[UserPublic])
def list_users(_: User = Depends(require_admin), db: Session = Depends(get_db)) -> List[UserPublic]:
    return [public_user(user) for user in db.scalars(select(User).order_by(User.created_at))]


@app.post("/api/v1/admin/users", response_model=UserPublic)
def create_user(payload: UserCreate, _: User = Depends(require_admin), db: Session = Depends(get_db)) -> UserPublic:
    if payload.role not in {"admin", "member"}:
        raise HTTPException(status_code=400, detail="角色不正确")
    if db.scalar(select(User).where(User.username == payload.username.strip())):
        raise HTTPException(status_code=409, detail="用户名已存在")
    user = User(
        username=payload.username.strip(),
        email=payload.email,
        display_name=payload.display_name.strip(),
        password_hash=hash_password(payload.password),
        role=payload.role,
        status="active",
    )
    db.add(user)
    bump_revision(db)
    db.commit()
    db.refresh(user)
    return public_user(user)


@app.patch("/api/v1/admin/users/{user_id}", response_model=UserPublic)
def update_user(user_id: str, payload: UserUpdate, _: User = Depends(require_admin), db: Session = Depends(get_db)) -> UserPublic:
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(status_code=404, detail="用户不存在")
    if payload.display_name is not None:
        target.display_name = payload.display_name.strip()
    if payload.status is not None:
        if payload.status not in {"active", "disabled"}:
            raise HTTPException(status_code=400, detail="账号状态不正确")
        target.status = payload.status
    bump_revision(db)
    db.commit()
    return public_user(target)


@app.post("/api/v1/admin/users/{user_id}/reset-password", response_model=ApiMessage)
def reset_password(user_id: str, payload: PasswordResetRequest, _: User = Depends(require_admin), db: Session = Depends(get_db)) -> ApiMessage:
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(status_code=404, detail="用户不存在")
    target.password_hash = hash_password(payload.password)
    db.query(AuthSession).filter(AuthSession.user_id == target.id).delete(synchronize_session=False)
    bump_revision(db)
    db.commit()
    return ApiMessage(message="密码已重置")


@app.get("/api/v1/companies")
def list_companies(user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> List[Dict[str, Any]]:
    query = select(Company).where(Company.status == "active", Company.owner_id == user.id).order_by(Company.name)
    return [{"id": item.id, "name": item.name, "status": item.status} for item in db.scalars(query)]


@app.post("/api/v1/companies")
def create_company(payload: CompanyCreate, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    company_name_value = payload.name.strip()
    normalized = normalize_name(company_name_value)
    if not company_name_value or not normalized:
        raise HTTPException(status_code=400, detail="公司名称不能为空")
    company = db.scalar(select(Company).where(Company.owner_id == user.id, Company.normalized_name == normalized))
    if company:
        return {"id": company.id, "name": company.name, "status": company.status, "existing": True}
    company = Company(owner_id=user.id, name=company_name_value, normalized_name=normalized)
    db.add(company)
    bump_revision(db)
    db.commit()
    db.refresh(company)
    return {"id": company.id, "name": company.name, "status": company.status, "existing": False}


@app.patch("/api/v1/companies/{company_id}")
def update_company(company_id: str, payload: CompanyUpdate, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    company = require_owned_company(db, company_id, user)
    if payload.name is not None:
        company_name_value = payload.name.strip()
        normalized = normalize_name(company_name_value)
        if not company_name_value or not normalized:
            raise HTTPException(status_code=400, detail="公司名称不能为空")
        duplicate = db.scalar(
            select(Company).where(
                Company.owner_id == user.id,
                Company.normalized_name == normalized,
                Company.id != company.id,
            )
        )
        if duplicate:
            raise HTTPException(status_code=409, detail="该公司名称已存在")
        company.name = company_name_value
        company.normalized_name = normalized
    if payload.status is not None:
        if payload.status not in {"active", "archived"}:
            raise HTTPException(status_code=400, detail="公司状态不正确")
        company.status = payload.status
    bump_revision(db)
    db.commit()
    return {"id": company.id, "name": company.name, "status": company.status}


def rule_payload(rule: ProjectTypeRule) -> Dict[str, Any]:
    return {
        "id": rule.id,
        "project_type_id": rule.project_type_id,
        "version_no": rule.version_no,
        "standard_performance_yuan": cents_to_yuan(rule.standard_performance_cents),
        "default_ratio_percent": bps_to_percent(rule.default_ratio_bps),
        "publicity_required": rule.publicity_required,
        "collection_stage": rule.collection_stage,
        "payout_pattern": rule.payout_pattern,
        "effective_from": rule.effective_from,
        "effective_to": rule.effective_to,
        "is_active": rule.is_active,
        "source_note": rule.source_note,
    }


def type_payload(db: Session, project_type: ProjectType) -> Dict[str, Any]:
    rules = list(db.scalars(select(ProjectTypeRule).where(ProjectTypeRule.project_type_id == project_type.id).order_by(ProjectTypeRule.version_no.desc())))
    return {
        "id": project_type.id,
        "name": project_type.name,
        "status": project_type.status,
        "created_at": project_type.created_at,
        "rules": [rule_payload(rule) for rule in rules],
    }


@app.get("/api/v1/project-types")
def list_project_types(_: User = Depends(get_current_user), db: Session = Depends(get_db)) -> List[Dict[str, Any]]:
    return [type_payload(db, item) for item in db.scalars(select(ProjectType).order_by(ProjectType.name))]


@app.post("/api/v1/project-types")
def create_project_type(payload: ProjectTypeCreate, user: User = Depends(require_admin), db: Session = Depends(get_db)) -> Dict[str, Any]:
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="项目类型名称不能为空")
    if db.scalar(select(ProjectType).where(ProjectType.name == name)):
        raise HTTPException(status_code=409, detail="项目类型已存在")
    project_type = ProjectType(name=name, created_by=user.id)
    db.add(project_type)
    bump_revision(db)
    db.commit()
    db.refresh(project_type)
    return type_payload(db, project_type)


@app.patch("/api/v1/project-types/{project_type_id}")
def update_project_type(project_type_id: str, payload: ProjectTypeUpdate, _: User = Depends(require_admin), db: Session = Depends(get_db)) -> Dict[str, Any]:
    project_type = db.get(ProjectType, project_type_id)
    if not project_type:
        raise HTTPException(status_code=404, detail="项目类型不存在")
    if payload.name is not None:
        name = payload.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="项目类型名称不能为空")
        duplicate = db.scalar(select(ProjectType).where(ProjectType.name == name, ProjectType.id != project_type.id))
        if duplicate:
            raise HTTPException(status_code=409, detail="项目类型已存在")
        project_type.name = name
    if payload.status is not None:
        if payload.status not in {"active", "disabled"}:
            raise HTTPException(status_code=400, detail="项目类型状态不正确")
        project_type.status = payload.status
    bump_revision(db)
    db.commit()
    return type_payload(db, project_type)


@app.post("/api/v1/project-types/{project_type_id}/rules")
def create_rule(project_type_id: str, payload: RuleCreate, _: User = Depends(require_admin), db: Session = Depends(get_db)) -> Dict[str, Any]:
    if not db.get(ProjectType, project_type_id):
        raise HTTPException(status_code=404, detail="项目类型不存在")
    if not is_supported_rule_combination(payload.collection_stage, payload.payout_pattern):
        raise HTTPException(status_code=400, detail="收款节点与发放模式组合不受支持")
    previous = db.scalar(select(func.max(ProjectTypeRule.version_no)).where(ProjectTypeRule.project_type_id == project_type_id)) or 0
    rule = ProjectTypeRule(
        project_type_id=project_type_id,
        version_no=previous + 1,
        standard_performance_cents=yuan_to_cents(payload.standard_performance_yuan) or 0,
        default_ratio_bps=percent_to_bps(payload.default_ratio_percent),
        publicity_required=payload.publicity_required,
        collection_stage=payload.collection_stage,
        payout_pattern=payload.payout_pattern,
        effective_from=payload.effective_from,
        effective_to=payload.effective_to,
        source_note=payload.source_note,
    )
    db.add(rule)
    bump_revision(db)
    db.commit()
    db.refresh(rule)
    return rule_payload(rule)


@app.get("/api/v1/records")
def list_records(
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    q: str = Query(default=""),
    current_status: str = Query(default=""),
    archived: bool = Query(default=False),
) -> List[Dict[str, Any]]:
    query = select(CooperationRecord).where(
        CooperationRecord.owner_id == user.id,
        CooperationRecord.archived_at.is_not(None) if archived else CooperationRecord.archived_at.is_(None),
    ).order_by(CooperationRecord.record_date.desc(), CooperationRecord.created_at.desc())
    records = list(db.scalars(query))
    search = normalize_name(q)
    output: List[Dict[str, Any]] = []
    for record in records:
        payload = record_payload(db, record)
        if search and search not in normalize_name(payload["company_name"] + payload["project_type_name"]):
            continue
        if current_status and payload["current_status"] != current_status:
            continue
        output.append(payload)
    return output


@app.post("/api/v1/records/duplicate-check")
def duplicate_check(payload: RecordCreate, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    normalized = normalize_name(payload.company_name)
    company = db.scalar(select(Company).where(Company.owner_id == user.id, Company.normalized_name == normalized))
    if not company:
        return {"duplicates": []}
    duplicates = list(db.scalars(
        select(CooperationRecord).where(
            CooperationRecord.owner_id == user.id,
            CooperationRecord.company_id == company.id,
            CooperationRecord.project_type_id == payload.project_type_id,
        ).order_by(CooperationRecord.record_date.desc())
    ))
    dates = sorted({record.record_date.isoformat() for record in duplicates}, reverse=True)
    date_note = f"，历史日期：{'、'.join(dates[:5])}" if dates else ""
    return {
        "duplicates": [record.id for record in duplicates],
        "message": f"发现相同公司、类型的历史记录{date_note}，确认后仍可新建" if duplicates else "",
    }


@app.post("/api/v1/records")
def create_record(payload: RecordCreate, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record_date = payload.record_date or today_beijing()
    project_type = db.get(ProjectType, payload.project_type_id)
    if not project_type or project_type.status != "active":
        raise HTTPException(status_code=400, detail="项目类型不存在或已停用")
    rule = rule_for_date(db, project_type.id, record_date)
    if not rule:
        raise HTTPException(status_code=400, detail="该项目类型在记录日期没有可用规则")
    if payload.participation_mode not in {"exclusive", "shared"}:
        raise HTTPException(status_code=400, detail="参与方式不正确")
    if payload.participation_mode == "exclusive" and payload.my_ratio_percent not in {None, 100}:
        raise HTTPException(status_code=400, detail="独享记录的个人比例固定为100%")
    if (payload.override_standard_yuan is not None or payload.manual_due_yuan is not None) and not payload.override_reason:
        raise HTTPException(status_code=400, detail="修改绩效金额时请填写调整原因")
    company_name_value = payload.company_name.strip()
    normalized = normalize_name(company_name_value)
    if not company_name_value or not normalized:
        raise HTTPException(status_code=400, detail="公司名称不能为空")
    company = db.scalar(select(Company).where(Company.owner_id == user.id, Company.normalized_name == normalized))
    if not company:
        company = Company(owner_id=user.id, name=company_name_value, normalized_name=normalized)
        db.add(company)
        db.flush()
    ratio = 10000 if payload.participation_mode == "exclusive" else percent_to_bps(payload.my_ratio_percent)
    record = CooperationRecord(
        owner_id=user.id,
        company_id=company.id,
        project_type_id=project_type.id,
        rule_id=rule.id,
        record_date=record_date,
        work_status="completed" if payload.completion_date else "in_progress",
        completion_date=payload.completion_date,
        publicity_status="pending" if rule.publicity_required else "not_applicable",
        participation_mode=payload.participation_mode,
        my_ratio_bps=ratio,
        standard_performance_cents=rule.standard_performance_cents,
        override_standard_cents=yuan_to_cents(payload.override_standard_yuan),
        manual_due_amount_cents=yuan_to_cents(payload.manual_due_yuan),
        snapshot_publicity_required=rule.publicity_required,
        snapshot_collection_stage=rule.collection_stage,
        snapshot_payout_pattern=rule.payout_pattern,
        note=payload.note,
        override_reason=payload.override_reason,
    )
    db.add(record)
    if payload.override_standard_yuan is not None or payload.manual_due_yuan is not None:
        db.flush()
        if payload.override_standard_yuan is not None:
            db.add(change_payload("standard", cents_to_yuan(record.standard_performance_cents), payload.override_standard_yuan, payload.override_reason or "", record))
        if payload.manual_due_yuan is not None:
            base_standard = record.override_standard_cents if record.override_standard_cents is not None else record.standard_performance_cents
            default_due = (base_standard * record.my_ratio_bps + 5000) // 10000 if base_standard is not None and record.my_ratio_bps is not None else None
            db.add(change_payload("due_amount", cents_to_yuan(default_due), payload.manual_due_yuan, payload.override_reason or "", record))
    bump_revision(db)
    db.commit()
    db.refresh(record)
    return record_payload(db, record)


@app.get("/api/v1/records/{record_id}")
def get_record(record_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    return record_payload(db, require_record(db, record_id, user))


@app.patch("/api/v1/records/{record_id}")
def update_record(record_id: str, payload: RecordUpdate, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    changes: List[PerformanceChange] = []
    if payload.company_name is not None:
        normalized = normalize_name(payload.company_name)
        company = db.scalar(select(Company).where(Company.owner_id == record.owner_id, Company.normalized_name == normalized))
        if not company:
            company = Company(owner_id=record.owner_id, name=payload.company_name.strip(), normalized_name=normalized)
            db.add(company)
            db.flush()
        record.company_id = company.id
    if payload.record_date is not None:
        record.record_date = payload.record_date
    supplied_fields = payload.model_fields_set
    if "completion_date" in supplied_fields or payload.work_status is not None:
        raise HTTPException(status_code=400, detail="请使用带原因的里程碑更正接口修改工作完成状态或日期")
    target_participation_mode = payload.participation_mode or record.participation_mode
    if target_participation_mode == "exclusive" and "my_ratio_percent" in supplied_fields and payload.my_ratio_percent != 100:
        raise HTTPException(status_code=400, detail="独享记录的个人比例固定为100%")
    if payload.participation_mode is not None:
        if payload.participation_mode not in {"exclusive", "shared"}:
            raise HTTPException(status_code=400, detail="参与方式不正确")
        if payload.participation_mode != record.participation_mode:
            if not payload.reason:
                raise HTTPException(status_code=400, detail="修改参与方式时请填写调整原因")
            changes.append(change_payload("participation_mode", record.participation_mode, payload.participation_mode, payload.reason, record))
        record.participation_mode = payload.participation_mode
        if payload.participation_mode == "exclusive" and record.my_ratio_bps != 10000:
            changes.append(change_payload("ratio", bps_to_percent(record.my_ratio_bps), 100, payload.reason or "独享默认100%", record))
            record.my_ratio_bps = 10000
    financial_updates = (
        ("my_ratio_percent", "ratio", lambda: bps_to_percent(record.my_ratio_bps), lambda value: setattr(record, "my_ratio_bps", percent_to_bps(value)), payload.my_ratio_percent),
        ("override_standard_yuan", "standard", lambda: cents_to_yuan(record.override_standard_cents), lambda value: setattr(record, "override_standard_cents", yuan_to_cents(value)), payload.override_standard_yuan),
        ("manual_due_yuan", "due_amount", lambda: cents_to_yuan(record.manual_due_amount_cents), lambda value: setattr(record, "manual_due_amount_cents", yuan_to_cents(value)), payload.manual_due_yuan),
    )
    changed_financial_values = []
    for request_field, history_field, get_old_value, set_value, new_value in financial_updates:
        if request_field not in supplied_fields:
            continue
        old_value = get_old_value()
        if old_value != new_value:
            changed_financial_values.append((history_field, old_value, new_value, set_value))
    if changed_financial_values and not payload.reason:
        raise HTTPException(status_code=400, detail="修改绩效金额或比例时请填写调整原因")
    for history_field, old_value, new_value, set_value in changed_financial_values:
        changes.append(change_payload(history_field, old_value, new_value, payload.reason or "", record))
        set_value(new_value)
    if changed_financial_values:
        has_manual_values = record.override_standard_cents is not None or record.manual_due_amount_cents is not None or (
            record.participation_mode == "shared" and record.my_ratio_bps is not None
        )
        record.override_reason = payload.reason if has_manual_values else None
    if payload.note is not None:
        record.note = payload.note
    for change in changes:
        db.add(change)
    bump_revision(db)
    db.commit()
    return record_payload(db, record)


@app.get("/api/v1/records/{record_id}/performance-changes")
def list_performance_changes(record_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> List[Dict[str, Any]]:
    record = require_record(db, record_id, user)
    changes = db.scalars(
        select(PerformanceChange)
        .where(PerformanceChange.owner_id == user.id, PerformanceChange.cooperation_record_id == record.id)
        .order_by(PerformanceChange.created_at.desc(), PerformanceChange.id.desc())
    )
    return [
        {
            "id": item.id,
            "field_name": item.field_name,
            "old_value": item.old_value,
            "new_value": item.new_value,
            "reason": item.reason,
            "created_at": item.created_at,
        }
        for item in changes
    ]


@app.post("/api/v1/records/{record_id}/complete")
def complete_record(record_id: str, payload: MilestoneRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    occurred_date = payload.occurred_date or today_beijing()
    if record.completion_date is not None and record.completion_date != occurred_date:
        raise HTTPException(status_code=400, detail="完成日期已记录；更正请填写原因并使用里程碑更正接口")
    record.work_status = "completed"
    record.completion_date = occurred_date
    if record.snapshot_publicity_required and record.publicity_status != "published":
        record.publicity_status = "pending"
    bump_revision(db)
    db.commit()
    return record_payload(db, record)


@app.post("/api/v1/records/{record_id}/publish")
def publish_record(record_id: str, payload: MilestoneRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    if not record.snapshot_publicity_required:
        raise HTTPException(status_code=400, detail="该项目不需要公示")
    occurred_date = payload.occurred_date or today_beijing()
    if record.publicity_date is not None and record.publicity_date != occurred_date:
        raise HTTPException(status_code=400, detail="公示日期已记录；更正请填写原因并使用里程碑更正接口")
    record.publicity_status = "published"
    record.publicity_date = occurred_date
    bump_revision(db)
    db.commit()
    return record_payload(db, record)


@app.post("/api/v1/records/{record_id}/advance-received")
def advance_received(record_id: str, payload: MilestoneRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    if record.snapshot_collection_stage not in {"advance", "advance_and_full"}:
        raise HTTPException(status_code=400, detail="该项目规则不包含预付款节点")
    occurred_date = payload.occurred_date or today_beijing()
    if record.advance_received_date is not None and record.advance_received_date != occurred_date:
        raise HTTPException(status_code=400, detail="预付款日期已记录；更正请填写原因并使用里程碑更正接口")
    record.advance_received_date = occurred_date
    bump_revision(db)
    db.commit()
    return record_payload(db, record)


@app.post("/api/v1/records/{record_id}/full-received")
def full_received(record_id: str, payload: MilestoneRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    if record.snapshot_collection_stage not in {"full", "advance_and_full"}:
        raise HTTPException(status_code=400, detail="该项目规则不包含全款节点")
    occurred_date = payload.occurred_date or today_beijing()
    if record.full_received_date is not None and record.full_received_date != occurred_date:
        raise HTTPException(status_code=400, detail="全款日期已记录；更正请填写原因并使用里程碑更正接口")
    record.full_received_date = occurred_date
    bump_revision(db)
    db.commit()
    return record_payload(db, record)


@app.patch("/api/v1/records/{record_id}/milestones")
def update_milestones(record_id: str, payload: MilestoneUpdateRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    supplied_fields = payload.model_fields_set - {"reason"}
    if not supplied_fields:
        raise HTTPException(status_code=400, detail="请选择至少一个里程碑日期")
    if "publicity_date" in supplied_fields and not record.snapshot_publicity_required and payload.publicity_date is not None:
        raise HTTPException(status_code=400, detail="该项目不需要公示")
    if "advance_received_date" in supplied_fields and record.snapshot_collection_stage not in {"advance", "advance_and_full"} and payload.advance_received_date is not None:
        raise HTTPException(status_code=400, detail="该项目规则不包含预付款节点")
    if "full_received_date" in supplied_fields and record.snapshot_collection_stage not in {"full", "advance_and_full"} and payload.full_received_date is not None:
        raise HTTPException(status_code=400, detail="该项目规则不包含全款节点")
    for field_name in supplied_fields:
        new_value = getattr(payload, field_name)
        old_value = getattr(record, field_name)
        if old_value == new_value:
            continue
        db.add(change_payload(field_name, old_value, new_value, payload.reason, record))
        setattr(record, field_name, new_value)
        if field_name == "completion_date":
            record.work_status = "completed" if new_value else "in_progress"
        elif field_name == "publicity_date":
            record.publicity_status = "published" if new_value else "pending"
    bump_revision(db)
    db.commit()
    return record_payload(db, record)


@app.post("/api/v1/records/{record_id}/override")
def override_record(record_id: str, payload: OverrideRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    supplied_fields = payload.model_fields_set - {"reason"}
    changes = [
        ("override_standard_yuan", "standard", cents_to_yuan(record.override_standard_cents), payload.override_standard_yuan),
        ("my_ratio_percent", "ratio", bps_to_percent(record.my_ratio_bps), payload.my_ratio_percent),
        ("manual_due_yuan", "due_amount", cents_to_yuan(record.manual_due_amount_cents), payload.manual_due_yuan),
    ]
    changed = False
    for request_field, field_name, old_value, new_value in changes:
        if request_field not in supplied_fields or old_value == new_value:
            continue
        db.add(change_payload(field_name, old_value, new_value, payload.reason, record))
        changed = True
        if request_field == "override_standard_yuan":
            record.override_standard_cents = yuan_to_cents(new_value)
        elif request_field == "my_ratio_percent":
            record.my_ratio_bps = percent_to_bps(new_value)
        else:
            record.manual_due_amount_cents = yuan_to_cents(new_value)
    if changed:
        has_manual_values = record.override_standard_cents is not None or record.manual_due_amount_cents is not None or (
            record.participation_mode == "shared" and record.my_ratio_bps is not None
        )
        record.override_reason = payload.reason if has_manual_values else None
    bump_revision(db)
    db.commit()
    return record_payload(db, record)


@app.post("/api/v1/records/{record_id}/archive", response_model=ApiMessage)
def archive_record(record_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> ApiMessage:
    record = require_record(db, record_id, user)
    record.archived_at = utc_now()
    bump_revision(db)
    db.commit()
    return ApiMessage(message="记录已归档")


@app.post("/api/v1/records/{record_id}/unarchive", response_model=ApiMessage)
def unarchive_record(record_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> ApiMessage:
    record = require_record(db, record_id, user)
    record.archived_at = None
    bump_revision(db)
    db.commit()
    return ApiMessage(message="记录已恢复")


@app.get("/api/v1/records/{record_id}/payouts")
def list_payouts(record_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> List[Dict[str, Any]]:
    record = require_record(db, record_id, user)
    return [payout_payload(item) for item in get_payouts(db, record.id)]


@app.post("/api/v1/records/{record_id}/payouts")
def create_payout(record_id: str, payload: PayoutCreate, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    record = require_record(db, record_id, user)
    if payload.batch_type not in {"advance", "final", "manual", "unknown"}:
        raise HTTPException(status_code=400, detail="到账批次不正确")
    payout = PayoutRecord(
        owner_id=record.owner_id,
        cooperation_record_id=record.id,
        batch_type=payload.batch_type,
        amount_cents=yuan_to_cents(payload.amount_yuan) or 0,
        received_date=payload.received_date or today_beijing(),
        note=payload.note,
    )
    db.add(payout)
    bump_revision(db)
    db.commit()
    db.refresh(payout)
    return payout_payload(payout)


@app.patch("/api/v1/payouts/{payout_id}")
def update_payout(payout_id: str, payload: PayoutUpdate, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    payout = db.get(PayoutRecord, payout_id)
    if not payout or payout.owner_id != user.id:
        raise HTTPException(status_code=404, detail="到账记录不存在")
    if payout.is_void:
        raise HTTPException(status_code=400, detail="已作废到账不能再次修改")
    record = require_record(db, payout.cooperation_record_id, user)
    supplied_fields = payload.model_fields_set - {"reason"}
    updates = {
        "amount_yuan": ("payout_amount", cents_to_yuan(payout.amount_cents), payload.amount_yuan),
        "received_date": ("payout_date", payout.received_date, payload.received_date),
        "batch_type": ("payout_batch", payout.batch_type, payload.batch_type),
        "note": ("payout_note", payout.note, payload.note),
    }
    for request_field in supplied_fields:
        if request_field not in updates:
            continue
        history_field, old_value, new_value = updates[request_field]
        if request_field == "received_date" and new_value is None:
            raise HTTPException(status_code=400, detail="到账日期不能清空")
        if request_field == "batch_type" and new_value not in {"advance", "final", "manual", "unknown"}:
            raise HTTPException(status_code=400, detail="到账批次不正确")
        if old_value == new_value:
            continue
        db.add(change_payload(history_field, old_value, new_value, payload.reason, record))
        if request_field == "amount_yuan":
            payout.amount_cents = yuan_to_cents(new_value) or 0
        elif request_field == "received_date":
            payout.received_date = new_value
        elif request_field == "batch_type":
            payout.batch_type = new_value
        else:
            payout.note = new_value
    bump_revision(db)
    db.commit()
    return payout_payload(payout)


@app.post("/api/v1/payouts/{payout_id}/void", response_model=ApiMessage)
def void_payout(payout_id: str, payload: PayoutVoidRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> ApiMessage:
    payout = db.get(PayoutRecord, payout_id)
    if not payout or payout.owner_id != user.id:
        raise HTTPException(status_code=404, detail="到账记录不存在")
    if payout.is_void:
        raise HTTPException(status_code=400, detail="到账已经作废")
    record = require_record(db, payout.cooperation_record_id, user)
    payout.is_void = True
    db.add(change_payload("payout_void", False, True, payload.reason, record))
    bump_revision(db)
    db.commit()
    return ApiMessage(message="到账记录已作废")


def scoped_records(db: Session, user: User) -> List[CooperationRecord]:
    query = select(CooperationRecord).where(
        CooperationRecord.owner_id == user.id,
        CooperationRecord.archived_at.is_(None),
    )
    return list(db.scalars(query))


@app.get("/api/v1/dashboard/summary", response_model=DashboardSummary)
def dashboard_summary(user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> DashboardSummary:
    records = scoped_records(db, user)
    status_counts: Dict[str, int] = {}
    paid_total = 0
    due_total = 0
    outstanding_total = 0
    unknown_due = 0
    completed_count = 0
    in_progress_count = 0
    for record in records:
        payouts = get_payouts(db, record.id)
        state = calculate_current_status(record, payouts)
        status_counts[state] = status_counts.get(state, 0) + 1
        if record.work_status == "completed":
            completed_count += 1
        else:
            in_progress_count += 1
        paid = calculate_paid_total(payouts)
        paid_total += paid
        due = calculate_due_amount(record)
        if due is None:
            unknown_due += 1
        else:
            due_total += due
            outstanding_total += max(due - paid, 0)
    return DashboardSummary(
        project_count=len(records),
        in_progress_count=in_progress_count,
        completed_count=completed_count,
        paid_total_cents=paid_total,
        due_total_cents=due_total,
        outstanding_total_cents=outstanding_total,
        unknown_due_count=unknown_due,
        status_counts=status_counts,
    )


HEADER_ALIASES = {
    "company_name": {"公司", "公司名", "公司名称", "合作公司", "客户", "合作方", "company", "companyname"},
    "project_type_name": {"类型", "项目类型", "项目类别", "项目类型名称", "type", "projecttype", "projecttypename"},
    "record_date": {"日期", "记录日期", "发生日期", "recorddate", "date"},
    "standard_performance_yuan": {"金额", "绩效", "标准绩效", "标准金额", "performance", "amount"},
    "my_ratio_percent": {"比例", "我的比例", "个人比例", "ratio", "myratio"},
}


def normalize_header(value: str) -> str:
    return re.sub(r"[\s_\-（）()]+", "", value.strip().casefold())


def parse_decimal(value: str) -> Optional[float]:
    if not value.strip():
        return None
    clean = value.strip().replace(",", "").replace("￥", "").replace("¥", "")
    try:
        parsed = Decimal(clean)
    except (InvalidOperation, ValueError):
        return None
    if not parsed.is_finite():
        return None
    converted = float(parsed)
    return converted if math.isfinite(converted) else None


def parse_ratio(value: str) -> Optional[float]:
    if not value.strip():
        return None
    text = value.strip().replace("％", "%")
    number = parse_decimal(text.replace("%", ""))
    if number is None:
        return None
    return number if "%" in text or number > 1 else number * 100


def split_table(content: str) -> Tuple[List[List[str]], str]:
    text = content.strip()
    if "\t" in text:
        delimiter = "\t"
    elif ";" in text and "," not in text:
        delimiter = ";"
    else:
        delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return rows, delimiter


def parse_rows(content: str) -> Tuple[List[Dict[str, str]], str, bool]:
    rows, delimiter = split_table(content)
    if not rows:
        return [], delimiter, False
    first_normalized = {normalize_header(value) for value in rows[0]}
    has_header = any(alias in first_normalized for aliases in HEADER_ALIASES.values() for alias in aliases)
    if has_header:
        index_map: Dict[str, int] = {}
        for index, header in enumerate(rows[0]):
            normalized = normalize_header(header)
            for field_name, aliases in HEADER_ALIASES.items():
                if normalized in {normalize_header(alias) for alias in aliases}:
                    index_map[field_name] = index
        data_rows = rows[1:]
    else:
        index_map = {"company_name": 0, "project_type_name": 1}
        if len(rows[0]) >= 3:
            index_map["record_date"] = 2
        if len(rows[0]) >= 4:
            index_map["standard_performance_yuan"] = 3
        if len(rows[0]) >= 5:
            index_map["my_ratio_percent"] = 4
        data_rows = rows
    output: List[Dict[str, str]] = []
    for row in data_rows:
        output.append({field: row[index].strip() if index < len(row) else "" for field, index in index_map.items()})
    return output, delimiter, has_header


def preview_import(db: Session, user: User, request: ImportRequest) -> ImportPreview:
    if len(request.content.encode("utf-8")) > MAX_IMPORT_FILE_BYTES:
        raise HTTPException(status_code=413, detail="导入内容不能超过 5 MB")
    raw_rows, delimiter, has_header = parse_rows(request.content)
    if len(raw_rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=413, detail=f"单次最多导入 {MAX_IMPORT_ROWS} 行")
    preview_rows: List[ImportRow] = []
    seen_import_pairs: Dict[Tuple[str, str], int] = {}
    for row_number, raw in enumerate(raw_rows, start=2 if has_header else 1):
        errors: List[str] = []
        warnings: List[str] = []
        company = raw.get("company_name", "").strip()
        type_name_value = raw.get("project_type_name", "").strip()
        if not company:
            errors.append("缺少公司名称")
        if not type_name_value:
            errors.append("缺少项目类型")
        project_type = db.scalar(select(ProjectType).where(ProjectType.name == type_name_value, ProjectType.status == "active")) if type_name_value else None
        if type_name_value and not project_type:
            errors.append("项目类型未匹配，请先在项目类型管理中创建")
        record_date = None
        if raw.get("record_date"):
            try:
                record_date = parse_date_value(raw["record_date"])
            except ValueError:
                errors.append("记录日期无法解析")
        if record_date is None:
            record_date = today_beijing()
        amount = parse_decimal(raw.get("standard_performance_yuan", ""))
        if raw.get("standard_performance_yuan") and amount is None:
            errors.append("绩效金额无法解析")
        elif amount is not None and amount < 0:
            errors.append("绩效金额不能小于 0")
        ratio = parse_ratio(raw.get("my_ratio_percent", ""))
        if raw.get("my_ratio_percent") and ratio is None:
            errors.append("个人比例无法解析")
        elif ratio is not None and not 0 <= ratio <= 100:
            errors.append("个人比例必须在 0% 到 100% 之间")
        if project_type and not errors:
            if not rule_for_date(db, project_type.id, record_date):
                errors.append("该记录日期没有匹配的项目类型绩效规则")
        if project_type and company and not errors:
            normalized = normalize_name(company)
            pair = (normalized, project_type.id)
            existing_company = db.scalar(select(Company).where(Company.owner_id == user.id, Company.normalized_name == normalized))
            if existing_company:
                existing_dates = list(db.scalars(
                    select(CooperationRecord.record_date)
                    .where(
                    CooperationRecord.owner_id == user.id,
                    CooperationRecord.company_id == existing_company.id,
                    CooperationRecord.project_type_id == project_type.id,
                    )
                    .order_by(CooperationRecord.record_date.desc())
                ))
                if existing_dates:
                    dates = "、".join(item.isoformat() for item in existing_dates[:5])
                    warnings.append(f"已有同公司、类型的历史记录（{dates}），如为再次合作可继续导入")
            first_row = seen_import_pairs.get(pair)
            if first_row is not None:
                warnings.append(f"与本批次第 {first_row} 行公司和类型相同，确认后仍可新增")
            else:
                seen_import_pairs[pair] = row_number
        preview_rows.append(
            ImportRow(
                row_number=row_number,
                company_name=company,
                project_type_name=type_name_value,
                record_date=record_date,
                standard_performance_yuan=amount,
                my_ratio_percent=ratio,
                matched_project_type_id=project_type.id if project_type else None,
                errors=errors,
                warnings=warnings,
            )
        )
    return ImportPreview(
        source_type=request.source_type,
        delimiter="TAB" if delimiter == "\t" else delimiter,
        has_header=has_header,
        rows=preview_rows,
        valid_count=sum(1 for item in preview_rows if not item.errors),
        warning_count=sum(1 for item in preview_rows if item.warnings),
        error_count=sum(1 for item in preview_rows if item.errors),
    )


@app.post("/api/v1/imports/preview", response_model=ImportPreview)
def import_preview(payload: ImportRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> ImportPreview:
    return preview_import(db, user, payload)


async def read_bounded_upload(file: UploadFile) -> bytes:
    content = await file.read(MAX_IMPORT_FILE_BYTES + 1)
    if len(content) > MAX_IMPORT_FILE_BYTES:
        raise HTTPException(status_code=413, detail="导入文件不能超过 5 MB")
    return content


def import_file_source_type(filename: str) -> str:
    lowered = filename.lower()
    if lowered.endswith(".xlsx"):
        return "xlsx"
    if lowered.endswith((".csv", ".txt")):
        return "csv"
    raise HTTPException(status_code=415, detail="仅支持 .xlsx、.csv 或 .txt 文件")


def read_upload_content(filename: str, content: bytes) -> str:
    if filename.lower().endswith(".xlsx"):
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_IMPORT_EXPANDED_BYTES:
                    raise HTTPException(status_code=413, detail="Excel 解压后的数据超过 25 MB")
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except HTTPException:
            raise
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise HTTPException(status_code=400, detail="Excel 文件无法解析") from exc
        rows = []
        try:
            for row_number, row in enumerate(workbook.active.iter_rows(values_only=True), start=1):
                if row_number > MAX_IMPORT_ROWS:
                    raise HTTPException(status_code=413, detail=f"单次最多导入 {MAX_IMPORT_ROWS} 行")
                values = []
                for value in row:
                    if isinstance(value, (datetime, date)):
                        values.append(value.isoformat()[:10])
                    else:
                        values.append("" if value is None else str(value))
                rows.append("\t".join(values))
        finally:
            workbook.close()
        return "\n".join(rows)
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="CSV/文本文件必须使用 UTF-8 编码") from exc


@app.post("/api/v1/imports/preview-file", response_model=ImportPreview)
async def import_preview_file(file: UploadFile = File(...), user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> ImportPreview:
    filename = file.filename or ""
    source_type = import_file_source_type(filename)
    content = await read_bounded_upload(file)
    text = read_upload_content(filename, content)
    return preview_import(db, user, ImportRequest(source_type=source_type, content=text))


@app.post("/api/v1/imports/commit-file")
async def import_commit_file(
    file: UploadFile = File(...),
    row_numbers: List[int] = Form(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> Dict[str, Any]:
    filename = file.filename or ""
    source_type = import_file_source_type(filename)
    content = await read_bounded_upload(file)
    request = ImportCommitRequest(source_type=source_type, content=read_upload_content(filename, content), row_numbers=row_numbers)
    return import_commit(request, user, db)


@app.post("/api/v1/imports/commit")
def import_commit(payload: ImportCommitRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> Dict[str, Any]:
    if len(set(payload.row_numbers)) != len(payload.row_numbers):
        raise HTTPException(status_code=400, detail="所选导入行不能重复")
    preview = preview_import(db, user, payload)
    selected = set(payload.row_numbers)
    preview_by_number = {row.row_number: row for row in preview.rows}
    unknown_numbers = selected - set(preview_by_number)
    if unknown_numbers:
        raise HTTPException(status_code=400, detail="所选行号不属于当前导入内容")
    selected_rows = [row for row in preview.rows if row.row_number in selected]
    invalid_rows = [row for row in selected_rows if row.errors]
    if invalid_rows:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "所选行中包含错误项，未写入任何记录",
                "rows": [{"row_number": row.row_number, "errors": row.errors} for row in invalid_rows],
            },
        )
    prepared_rows: List[Tuple[ImportRow, ProjectType, ProjectTypeRule]] = []
    for preview_row in selected_rows:
        project_type = db.get(ProjectType, preview_row.matched_project_type_id)
        if not project_type or project_type.status != "active":
            raise HTTPException(status_code=409, detail=f"第 {preview_row.row_number} 行的项目类型已不可用，请重新预览")
        record_date = preview_row.record_date or today_beijing()
        rule = rule_for_date(db, project_type.id, record_date)
        if not rule:
            raise HTTPException(status_code=409, detail=f"第 {preview_row.row_number} 行的绩效规则已变化，请重新预览")
        prepared_rows.append((preview_row, project_type, rule))

    created_ids: List[str] = []
    created_row_numbers: List[int] = []
    skipped = [
        {
            "row_number": row.row_number,
            "reason": "用户未选择",
            "errors": row.errors,
            "warnings": row.warnings,
        }
        for row in preview.rows
        if row.row_number not in selected
    ]
    warnings = [
        {"row_number": row.row_number, "warnings": row.warnings}
        for row in selected_rows
        if row.warnings
    ]
    batch = ImportBatch(
        owner_id=user.id,
        source_type=payload.source_type,
        parser_mode="deterministic",
        status="committed",
    )
    try:
        for preview_row, project_type, rule in prepared_rows:
            record_date = preview_row.record_date or today_beijing()
            company_name_value = preview_row.company_name
            company = db.scalar(
                select(Company).where(
                    Company.owner_id == user.id,
                    Company.normalized_name == normalize_name(company_name_value),
                )
            )
            if not company:
                company = Company(owner_id=user.id, name=company_name_value, normalized_name=normalize_name(company_name_value))
                db.add(company)
                db.flush()
            ratio = percent_to_bps(preview_row.my_ratio_percent)
            record = CooperationRecord(
                owner_id=user.id,
                company_id=company.id,
                project_type_id=project_type.id,
                rule_id=rule.id,
                record_date=record_date,
                work_status="in_progress",
                publicity_status="pending" if rule.publicity_required else "not_applicable",
                participation_mode="shared" if ratio is not None else "exclusive",
                my_ratio_bps=ratio if ratio is not None else 10000,
                standard_performance_cents=yuan_to_cents(preview_row.standard_performance_yuan) if preview_row.standard_performance_yuan is not None else rule.standard_performance_cents,
                snapshot_publicity_required=rule.publicity_required,
                snapshot_collection_stage=rule.collection_stage,
                snapshot_payout_pattern=rule.payout_pattern,
            )
            db.add(record)
            db.flush()
            created_ids.append(record.id)
            created_row_numbers.append(preview_row.row_number)
        batch.created_record_ids_json = json.dumps(created_ids)
        batch.error_summary_json = json.dumps({"skipped": skipped, "warnings": warnings}, ensure_ascii=False)
        db.add(batch)
        bump_revision(db)
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail="导入写入失败，已整批回滚") from exc
    return {
        "batch_id": batch.id,
        "created_ids": created_ids,
        "created_row_numbers": created_row_numbers,
        "skipped": skipped,
        "warnings": warnings,
    }


@app.get("/api/v1/exports/records.csv")
def export_records(user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> StreamingResponse:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["合作公司", "项目类型", "记录日期", "工作状态", "当前节点", "标准绩效", "应得绩效", "已到账", "待到账"])
    for record in scoped_records(db, user):
        payload = record_payload(db, record)
        writer.writerow(
            [
                payload["company_name"],
                payload["project_type_name"],
                payload["record_date"],
                payload["work_status"],
                payload["current_status_label"],
                payload["standard_performance_yuan"],
                payload["my_due_amount_yuan"],
                payload["paid_total_yuan"],
                payload["outstanding_yuan"],
            ]
        )
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue().encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=records.csv"},
    )


@app.get("/api/v1/exports/records.json")
def export_records_json(user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> JSONResponse:
    return JSONResponse(content=jsonable_encoder({"exported_at": utc_now().isoformat(), "records": [record_payload(db, record) for record in scoped_records(db, user)]}))
